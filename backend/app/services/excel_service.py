import io
import re
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelService:
    @staticmethod
    def generate_template(for_story: bool = False) -> io.BytesIO:
        wb = openpyxl.Workbook()
        
        # Sheet 1: Casos de prueba
        ws_cases = wb.active
        ws_cases.title = 'Casos de prueba'
        ws_cases.views.sheetView[0].showGridLines = True
        
        if for_story:
            headers = [
                ('Código', 15),
                ('Título', 35),
                ('Módulo', 20),
                ('Descripción', 35),
                ('Precondiciones', 30),
                ('Pasos a reproducir', 45),
                ('Resultado esperado', 45),
                ('Prioridad', 15),
            ]
        else:
            headers = [
                ('Código', 15),
                ('Título', 35),
                ('Módulo', 20),
                ('Precondiciones', 30),
                ('Pasos', 45),
                ('Resultado esperado', 45),
                ('Prioridad', 15),
            ]
        
        header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='0F766E', end_color='0F766E', fill_type='solid')
        thin_border = Border(
            left=Side(style='thin', color='CCCCCC'),
            right=Side(style='thin', color='CCCCCC'),
            top=Side(style='thin', color='CCCCCC'),
            bottom=Side(style='thin', color='CCCCCC')
        )
        align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)

        for col_num, (header_text, width) in enumerate(headers, 1):
            cell = ws_cases.cell(row=1, column=col_num, value=header_text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = thin_border
            col_letter = get_column_letter(col_num)
            ws_cases.column_dimensions[col_letter].width = width

        ws_cases.freeze_panes = 'A2'
        ws_cases.row_dimensions[1].height = 28

        # Example row
        if for_story:
            example_row = [
                'WF-001',
                'Crear nuevo proceso BPMN',
                'Procesos',
                'Permite diseñar y guardar un nuevo flujo BPMN en el sistema.',
                'Usuario autenticado con permisos de administrador',
                '1. Entrar al módulo Procesos\n2. Pulsar Nuevo proceso\n3. Completar los datos\n4. Guardar',
                'El proceso se crea correctamente y aparece en el catálogo.',
                'P1'
            ]
        else:
            example_row = [
                'WF-001',
                'Crear nuevo proceso BPMN',
                'Procesos',
                'Usuario autenticado con permisos de administrador',
                '1. Entrar al módulo Procesos\n2. Pulsar Nuevo proceso\n3. Completar los datos\n4. Guardar',
                'El proceso se crea correctamente y aparece en el catálogo.',
                'P1'
            ]
        
        for col_num, val in enumerate(example_row, 1):
            cell = ws_cases.cell(row=2, column=col_num, value=val)
            cell.alignment = align_left if col_num not in (1, len(example_row)) else align_center
            cell.border = thin_border
            cell.font = Font(name='Calibri', size=10)
        ws_cases.row_dimensions[2].height = 65

        # Sheet 2: Instrucciones
        ws_info = wb.create_sheet(title='Instrucciones')
        ws_info.views.sheetView[0].showGridLines = True
        
        ws_info.column_dimensions['A'].width = 25
        ws_info.column_dimensions['B'].width = 65

        title_cell = ws_info.cell(row=1, column=1, value='Instrucciones para la importación de Casos de Prueba')
        title_cell.font = Font(name='Calibri', size=14, bold=True, color='0F766E')
        ws_info.row_dimensions[1].height = 26

        if for_story:
            instructions = [
                ('Campos obligatorios', 'Título, Resultado esperado y Prioridad son campos obligatorios. Código es opcional (si se omite, se generará automáticamente).'),
                ('Prioridades válidas', 'Los valores permitidos son: P1 (Crítico/Alto), P2 (Medio), P3 (Bajo), P4 (Opcional).'),
                ('Formato de Pasos', 'Escriba los pasos separados por saltos de línea (Alt + Enter) o numerados como 1. ..., 2. ...'),
                ('Descripción y Precondiciones', 'Campos de texto opcionales sobre el alcance y condiciones previas para ejecutar la prueba.'),
                ('Módulo', 'Especifique el nombre del módulo (ej. Procesos, Catálogo, FUID, Workflow). Si no existe, se vinculará al proyecto general.'),
                ('Códigos duplicados', 'El importador permite elegir si omitir duplicados o actualizar los casos existentes en el sistema.'),
                ('Trazabilidad y seguridad', 'Actualizar casos existentes preserva las ejecuciones históricas y evidencias sin alterarlas.')
            ]
        else:
            instructions = [
                ('Campos obligatorios', 'Código, Título, Resultado esperado y Prioridad son campos obligatorios.'),
                ('Prioridades válidas', 'Los valores permitidos son: P1 (Crítico/Alto), P2 (Medio), P3 (Bajo), P4 (Opcional).'),
                ('Formato de Pasos', 'Escriba los pasos separados por saltos de línea (Alt + Enter) o numerados como 1. ..., 2. ...'),
                ('Módulo', 'Especifique el nombre del módulo (ej. Procesos, Catálogo, FUID, Workflow). Si no existe, se vinculará al proyecto general.'),
                ('Códigos duplicados', 'El importador permite elegir si omitir duplicados o actualizar los casos existentes en el sistema.'),
                ('Trazabilidad y seguridad', 'Actualizar casos existentes preserva las ejecuciones históricas y evidencias sin alterarlas.')
            ]

        subhead_fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
        for idx, (label, desc) in enumerate(instructions, 3):
            c1 = ws_info.cell(row=idx, column=1, value=label)
            c1.font = Font(name='Calibri', size=11, bold=True, color='1E293B')
            c1.fill = subhead_fill
            c1.border = thin_border
            c1.alignment = Alignment(vertical='top')

            c2 = ws_info.cell(row=idx, column=2, value=desc)
            c2.font = Font(name='Calibri', size=11)
            c2.border = thin_border
            c2.alignment = Alignment(vertical='top', wrap_text=True)
            ws_info.row_dimensions[idx].height = 34

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def parse_and_validate_excel(content: bytes, existing_codes: set, auto_generate_code: bool = False) -> List[Dict[str, Any]]:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        if 'Casos de prueba' in wb.sheetnames:
            ws = wb['Casos de prueba']
        else:
            ws = wb.active

        results = []
        excel_seen_codes = set()
        valid_priorities = {'P1', 'P2', 'P3', 'P4'}

        header_map = {}
        for col_idx in range(1, ws.max_column + 1):
            val = ws.cell(row=1, column=col_idx).value
            if val:
                header_clean = str(val).strip().lower()
                if 'código' in header_clean or 'codigo' in header_clean or 'code' in header_clean:
                    header_map['code'] = col_idx
                elif 'título' in header_clean or 'titulo' in header_clean or 'caso' in header_clean or 'title' in header_clean:
                    header_map['title'] = col_idx
                elif 'módulo' in header_clean or 'modulo' in header_clean or 'module' in header_clean:
                    header_map['module'] = col_idx
                elif 'descripci' in header_clean or 'summary' in header_clean or 'descripcion' in header_clean:
                    header_map['description'] = col_idx
                elif 'precondici' in header_clean:
                    header_map['preconditions'] = col_idx
                elif 'paso' in header_clean or 'step' in header_clean:
                    header_map['steps'] = col_idx
                elif 'resultado' in header_clean or 'expected' in header_clean:
                    header_map['expected_result'] = col_idx
                elif 'prioridad' in header_clean or 'priority' in header_clean:
                    header_map['priority'] = col_idx

        if 'code' not in header_map: header_map['code'] = 1
        if 'title' not in header_map: header_map['title'] = 2
        if 'module' not in header_map: header_map['module'] = 3
        # If description was found, use it; otherwise check if col 4 is available and not preconditions
        if 'description' not in header_map and 'preconditions' in header_map and header_map['preconditions'] > 4:
            header_map['description'] = 4
        if 'preconditions' not in header_map: header_map['preconditions'] = 5 if 'description' in header_map else 4
        if 'steps' not in header_map: header_map['steps'] = 6 if 'description' in header_map else 5
        if 'expected_result' not in header_map: header_map['expected_result'] = 7 if 'description' in header_map else 6
        if 'priority' not in header_map: header_map['priority'] = 8 if 'description' in header_map else 7

        # Collect raw non-empty rows first
        raw_rows = []
        for row_idx in range(2, ws.max_row + 1):
            row_vals = [
                ws.cell(row=row_idx, column=header_map.get('code', 1)).value,
                ws.cell(row=row_idx, column=header_map.get('title', 2)).value,
                ws.cell(row=row_idx, column=header_map.get('expected_result', 6)).value,
                ws.cell(row=row_idx, column=header_map.get('priority', 7)).value,
            ]
            if any(v is not None and str(v).strip() != '' for v in row_vals):
                raw_rows.append(row_idx)

        if len(raw_rows) > 500:
            raise ValueError(f"El archivo supera el límite máximo de 500 casos de prueba ({len(raw_rows)} filas encontradas).")

        auto_code_counter = 1
        for c in existing_codes:
            if c:
                m = re.search(r"TC-(\d+)", c)
                if m:
                    auto_code_counter = max(auto_code_counter, int(m.group(1)) + 1)
        for row_idx in raw_rows:
            code_val = ws.cell(row=row_idx, column=header_map['code']).value
            title_val = ws.cell(row=row_idx, column=header_map['title']).value
            module_val = ws.cell(row=row_idx, column=header_map.get('module', 3)).value if 'module' in header_map else None
            desc_val = ws.cell(row=row_idx, column=header_map['description']).value if 'description' in header_map else None
            precond_val = ws.cell(row=row_idx, column=header_map.get('preconditions', 4)).value if 'preconditions' in header_map else None
            steps_val = ws.cell(row=row_idx, column=header_map.get('steps', 5)).value if 'steps' in header_map else None
            expected_val = ws.cell(row=row_idx, column=header_map.get('expected_result', 6)).value if 'expected_result' in header_map else None
            priority_val = ws.cell(row=row_idx, column=header_map.get('priority', 7)).value if 'priority' in header_map else None

            code_str = str(code_val).strip() if code_val is not None else ''
            title_str = str(title_val).strip() if title_val is not None else ''
            module_str = str(module_val).strip() if module_val is not None else ''
            desc_str = str(desc_val).strip() if desc_val is not None else ''
            precond_str = str(precond_val).strip() if precond_val is not None else ''
            expected_str = str(expected_val).strip() if expected_val is not None else ''
            priority_str = str(priority_val).strip().upper() if priority_val is not None else 'P2'

            steps_list = []
            if steps_val is not None:
                lines = [s.strip() for s in str(steps_val).split('\n') if s.strip()]
                steps_list = lines

            errors = []
            warnings = []

            # If code is omitted, auto-generate non-conflicting code
            if not code_str:
                if auto_generate_code:
                    while True:
                        cand = f"TC-{auto_code_counter:03d}"
                        auto_code_counter += 1
                        if cand not in existing_codes and cand not in excel_seen_codes:
                            code_str = cand
                            excel_seen_codes.add(code_str)
                            break
                else:
                    errors.append(f'Fila {row_idx}: El campo Código es obligatorio.')
            else:
                if len(code_str) > 50:
                    errors.append(f'Fila {row_idx}: El código supera el límite máximo de 50 caracteres.')
                if code_str in excel_seen_codes:
                    errors.append(f'Fila {row_idx}: Código duplicado en el archivo Excel ({code_str}).')
                else:
                    excel_seen_codes.add(code_str)

            if not title_str:
                errors.append(f'Fila {row_idx}: El campo Título es obligatorio.')
            elif len(title_str) > 255:
                errors.append(f'Fila {row_idx}: El título supera el límite máximo de 255 caracteres.')

            if not expected_str:
                errors.append(f'Fila {row_idx}: El Resultado esperado es obligatorio.')

            if priority_str not in valid_priorities:
                errors.append(f'Fila {row_idx}: La prioridad \'{priority_str}\' no es válida. Valores permitidos: P1, P2, P3, P4.')

            exists_in_db = code_str in existing_codes
            if exists_in_db and not errors:
                warnings.append(f'Fila {row_idx}: El código {code_str} ya existe en la base de datos.')

            status = 'ERROR' if errors else ('WARNING' if warnings else 'VALID')
            message = ' · '.join(errors + warnings) if (errors or warnings) else 'Válido'

            results.append({
                'row_number': row_idx,
                'code': code_str,
                'title': title_str,
                'module_name': module_str or None,
                'description': desc_str or None,
                'preconditions': precond_str or None,
                'steps': steps_list,
                'expected_result': expected_str,
                'priority': priority_str,
                'status': status,
                'message': message,
                'exists_in_db': exists_in_db
            })

        return results

    @staticmethod
    def generate_error_report(errors: List[Dict[str, Any]]) -> io.BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Errores de importación'
        ws.views.sheetView[0].showGridLines = True

        headers = [
            ('Fila', 10),
            ('Código', 15),
            ('Título', 30),
            ('Módulo', 20),
            ('Prioridad', 12),
            ('Error reportado', 55)
        ]

        header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='B91C1C', end_color='B91C1C', fill_type='solid')
        thin_border = Border(
            left=Side(style='thin', color='CCCCCC'),
            right=Side(style='thin', color='CCCCCC'),
            top=Side(style='thin', color='CCCCCC'),
            bottom=Side(style='thin', color='CCCCCC')
        )

        for col_num, (h_text, width) in enumerate(headers, 1):
            c = ws.cell(row=1, column=col_num, value=h_text)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = thin_border
            ws.column_dimensions[get_column_letter(col_num)].width = width
        ws.row_dimensions[1].height = 26

        for idx, err in enumerate(errors, 2):
            ws.cell(row=idx, column=1, value=err.get('row_number', '')).alignment = Alignment(horizontal='center')
            ws.cell(row=idx, column=2, value=err.get('code', '')).alignment = Alignment(horizontal='center')
            ws.cell(row=idx, column=3, value=err.get('title', ''))
            ws.cell(row=idx, column=4, value=err.get('module_name', ''))
            ws.cell(row=idx, column=5, value=err.get('priority', '')).alignment = Alignment(horizontal='center')
            ws.cell(row=idx, column=6, value=err.get('message', ''))
            for col_num in range(1, 7):
                ws.cell(row=idx, column=col_num).border = thin_border
            ws.row_dimensions[idx].height = 22

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
