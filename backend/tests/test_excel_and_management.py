import io
import pytest
from fastapi.testclient import TestClient
import openpyxl
from app.main import app

client = TestClient(app)

def get_auth_header():
    resp = client.post("/api/v1/auth/login", json={"email": "qa@gestadoc.com", "password": "qa123456"})
    token = resp.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}

def test_excel_template_generation():
    headers = get_auth_header()
    resp = client.get("/api/v1/test-cases/template", headers=headers)
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert "Casos de prueba" in wb.sheetnames
    assert "Instrucciones" in wb.sheetnames
    ws = wb["Casos de prueba"]
    headers_row = [ws.cell(row=1, column=i).value for i in range(1, 8)]
    assert headers_row == ["Código", "Título", "Módulo", "Precondiciones", "Pasos", "Resultado esperado", "Prioridad"]

def test_excel_import_preview_and_commit():
    headers = get_auth_header()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Casos de prueba"
    ws.append(["Código", "Título", "Módulo", "Precondiciones", "Pasos", "Resultado esperado", "Prioridad"])
    ws.append(["TEST-IMP-01", "Caso Válido 1", "Procesos", "Precondición", "1. Abrir\n2. Guardar", "Éxito", "P1"])
    ws.append(["TEST-IMP-02", "Caso con Warning Dup", "Procesos", "", "Paso 1", "Éxito", "P2"])
    ws.append(["", "Caso sin código (Error)", "Procesos", "", "", "Éxito", "P2"])
    ws.append(["TEST-IMP-03", "Caso prioridad inválida", "Procesos", "", "", "Éxito", "URGENTE"])

    file_bytes = io.BytesIO()
    wb.save(file_bytes)
    file_bytes.seek(0)

    files = {"file": ("test_cases.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    prev_resp = client.post("/api/v1/test-cases/import/preview?project_id=1", headers=headers, files=files)
    assert prev_resp.status_code == 200
    data = prev_resp.json()
    assert data["total_rows"] == 4
    assert data["valid_count"] >= 1
    assert data["error_count"] == 2

    import time
    stamp = int(time.time() * 1000)
    test_code = f"TEST-IMP-{stamp}"

    valid_items = [{
        "row_number": 2,
        "code": test_code,
        "title": "Caso Válido 1",
        "module_name": "Procesos",
        "preconditions": "Precondición",
        "steps": ["1. Abrir", "2. Guardar"],
        "expected_result": "Éxito",
        "priority": "P1"
    }]
    commit_payload = {"project_id": 1, "strategy": "SKIP", "items": valid_items}
    commit_resp = client.post("/api/v1/test-cases/import/commit", headers=headers, json=commit_payload)
    assert commit_resp.status_code == 200
    cdata = commit_resp.json()
    assert cdata["processed"] == 1
    assert cdata["created"] == 1

def test_duplicate_strategy_update():
    headers = get_auth_header()
    import time
    stamp = int(time.time() * 1000)
    test_code = f"TEST-UPD-{stamp}"
    
    # First create it
    create_payload = {"project_id": 1, "strategy": "SKIP", "items": [{
        "row_number": 2,
        "code": test_code,
        "title": "Caso Inicial",
        "module_name": "Procesos",
        "expected_result": "Inicial",
        "priority": "P1"
    }]}
    c1 = client.post("/api/v1/test-cases/import/commit", headers=headers, json=create_payload)
    assert c1.status_code == 200

    # Then update it via UPDATE strategy
    update_items = [{
        "row_number": 2,
        "code": test_code,
        "title": "Caso Válido 1 Modificado",
        "module_name": "Workflow",
        "preconditions": "Nueva precondicion",
        "steps": ["Paso modificado"],
        "expected_result": "Resultado modificado",
        "priority": "P3"
    }]
    commit_payload = {"project_id": 1, "strategy": "UPDATE", "items": update_items}
    commit_resp = client.post("/api/v1/test-cases/import/commit", headers=headers, json=commit_payload)
    assert commit_resp.status_code == 200
    cdata = commit_resp.json()
    assert cdata["updated"] == 1

def test_delete_and_safe_archive_behavior():
    headers = get_auth_header()
    import time
    stamp = int(time.time() * 1000)
    unique_code = f"TEMP-DEL-{stamp}"
    new_tc = client.post("/api/v1/test-cases", headers=headers, json={
        "project_id": 1,
        "code": unique_code,
        "title": "Caso para borrar sin ejecuciones",
        "expected_result": "Borrado",
        "priority": "P2"
    }).json()
    tc_id = new_tc["id"]

    del_resp = client.delete(f"/api/v1/test-cases/{tc_id}", headers=headers)
    assert del_resp.status_code == 200
    assert del_resp.json()["archived"] is False

    # Create a test case and attach a test result execution to it to verify soft-archive
    exec_tc = client.post("/api/v1/test-cases", headers=headers, json={
        "project_id": 1,
        "code": f"TEMP-ARC-{stamp}",
        "title": "Caso con ejecucion para archivar",
        "expected_result": "Resultado",
        "priority": "P1"
    }).json()
    exec_tc_id = exec_tc["id"]

    # Submit a test result for it using form-data expected by executions endpoint
    res_resp = client.post("/api/v1/executions/1/results", headers=headers, data={
        "test_case_id": exec_tc_id,
        "status": "PASSED",
        "actual_result": "Ejecutado con exito"
    })
    assert res_resp.status_code == 200

    del_arc = client.delete(f"/api/v1/test-cases/{exec_tc_id}", headers=headers)
    assert del_arc.status_code == 200
    assert del_arc.json()["archived"] is True