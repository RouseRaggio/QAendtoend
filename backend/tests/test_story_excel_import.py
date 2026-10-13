import io
import time
import openpyxl
import pytest
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)

def get_auth_header():
    resp = client.post("/api/v1/auth/login", json={"email": "qa@gestadoc.com", "password": "qa123456"})
    token = resp.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}

def test_download_story_import_template():
    headers = get_auth_header()
    # Pick first available user story
    stories_resp = client.get("/api/v1/user-stories?project_id=1", headers=headers)
    assert stories_resp.status_code == 200
    stories = stories_resp.json()
    assert len(stories) > 0
    story_id = stories[0]["id"]

    resp = client.get(f"/api/v1/user-stories/{story_id}/import-cases/template", headers=headers)
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert "Casos de prueba" in wb.sheetnames
    assert "Instrucciones" in wb.sheetnames

    ws = wb["Casos de prueba"]
    headers_row = [ws.cell(row=1, column=i).value for i in range(1, 9)]
    assert headers_row == [
        "Código",
        "Título",
        "Módulo",
        "Descripción",
        "Precondiciones",
        "Pasos a reproducir",
        "Resultado esperado",
        "Prioridad"
    ]

    # Non-existent story returns 404
    bad_resp = client.get("/api/v1/user-stories/999999/import-cases/template", headers=headers)
    assert bad_resp.status_code == 404

def test_story_excel_import_preview_with_auto_code():
    headers = get_auth_header()
    stories_resp = client.get("/api/v1/user-stories?project_id=1", headers=headers)
    assert stories_resp.status_code == 200
    story = stories_resp.json()[0]
    story_id = story["id"]
    project_id = story["project_id"]

    stamp = int(time.time() * 1000)
    explicit_code = f"STORY-TC-{stamp}"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Casos de prueba"
    ws.append([
        "Código", "Título", "Módulo", "Descripción",
        "Precondiciones", "Pasos a reproducir", "Resultado esperado", "Prioridad"
    ])
    # Row 1: Valid with explicit code
    ws.append([
        explicit_code, "Caso explícito para historia", "Flujos", "Descripción detallada",
        "Estar logueado", "1. Paso A\n2. Paso B", "Todo ok", "P1"
    ])
    # Row 2: Empty code (should be auto-generated as TC-xxx)
    ws.append([
        "", "Caso con código omitido", "", "Sin código provisto",
        "Condición previa", "1. Paso auto", "Se auto-genera código", "P2"
    ])
    # Row 3: Invalid priority
    ws.append([
        f"STORY-BAD-{stamp}", "Caso con prioridad inválida", "", "",
        "", "1. Paso", "Error esperado", "URGENTISIMO"
    ])

    file_bytes = io.BytesIO()
    wb.save(file_bytes)
    file_bytes.seek(0)

    files = {"file": ("story_import.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    prev_resp = client.post(
        f"/api/v1/user-stories/{story_id}/import-cases/preview?project_id={project_id}",
        headers=headers,
        files=files
    )
    assert prev_resp.status_code == 200
    data = prev_resp.json()
    assert data["total_rows"] == 3
    assert data["valid_count"] == 2
    assert data["error_count"] == 1

    # Check auto-generated code on row 2
    item_auto = data["items"][1]
    assert item_auto["code"].startswith("TC-")
    assert item_auto["status"] == "VALID"
    assert item_auto["title"] == "Caso con código omitido"

    # Check error on row 3
    item_err = data["items"][2]
    assert item_err["status"] == "ERROR"
    assert "prioridad" in item_err["message"].lower()

def test_story_excel_import_preview_cross_project_rejected():
    headers = get_auth_header()
    stories_resp = client.get("/api/v1/user-stories?project_id=1", headers=headers)
    story = stories_resp.json()[0]
    story_id = story["id"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Código", "Título", "Resultado esperado", "Prioridad"])
    ws.append(["CODE-X", "Título X", "Resultado X", "P1"])
    file_bytes = io.BytesIO()
    wb.save(file_bytes)
    file_bytes.seek(0)

    files = {"file": ("test.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    resp = client.post(
        f"/api/v1/user-stories/{story_id}/import-cases/preview?project_id=9999",
        headers=headers,
        files=files
    )
    assert resp.status_code == 400
    assert "No se permite asociación cruzada" in resp.json()["detail"]

def test_story_excel_import_commit_and_auto_association():
    headers = get_auth_header()
    # Create or use a dedicated user story
    create_story_resp = client.post(
        "/api/v1/user-stories",
        headers=headers,
        json={
            "project_id": 1,
            "title": f"Historia para Import Excel {int(time.time())}",
            "description": "Validación de importación directa a HU",
            "status": "TODO"
        }
    )
    assert create_story_resp.status_code == 201
    story = create_story_resp.json()
    story_id = story["id"]
    project_id = story["project_id"]

    stamp = int(time.time() * 1000)
    tc1_code = f"HU-TC-{stamp}-1"
    tc2_code = f"HU-TC-{stamp}-2"

    items_to_commit = [
        {
            "row_number": 2,
            "code": tc1_code,
            "title": "Prueba de login directo",
            "module_name": "Seguridad",
            "description": "Descripción caso 1",
            "preconditions": "Precondición 1",
            "steps": ["Paso 1", "Paso 2"],
            "expected_result": "Login correcto",
            "priority": "P1"
        },
        {
            "row_number": 3,
            "code": tc2_code,
            "title": "Prueba de recuperación",
            "module_name": "Seguridad",
            "description": "Descripción caso 2",
            "preconditions": "Precondición 2",
            "steps": ["Paso 1", "Paso 2"],
            "expected_result": "Email enviado",
            "priority": "P2"
        }
    ]

    commit_resp = client.post(
        f"/api/v1/user-stories/{story_id}/import-cases/commit",
        headers=headers,
        json={
            "project_id": project_id,
            "strategy": "SKIP",
            "items": items_to_commit
        }
    )
    assert commit_resp.status_code == 200
    res_data = commit_resp.json()
    assert res_data["created"] == 2
    assert res_data["total_associated_cases"] == 2

    # Verify story detail returns both test cases associated
    detail_resp = client.get(f"/api/v1/user-stories/{story_id}", headers=headers)
    assert detail_resp.status_code == 200
    detail_data = detail_resp.json()
    assert len(detail_data["test_cases"]) == 2
    story_case_codes = [c["code"] for c in detail_data["test_cases"]]
    assert tc1_code in story_case_codes
    assert tc2_code in story_case_codes
    assert detail_data["test_cases_count"] == 2

def test_story_excel_import_commit_cross_project_rejected():
    headers = get_auth_header()
    stories_resp = client.get("/api/v1/user-stories?project_id=1", headers=headers)
    story = stories_resp.json()[0]
    story_id = story["id"]

    commit_resp = client.post(
        f"/api/v1/user-stories/{story_id}/import-cases/commit",
        headers=headers,
        json={
            "project_id": 9999,
            "strategy": "SKIP",
            "items": []
        }
    )
    assert commit_resp.status_code == 400
    assert "No se permite asociación cruzada" in commit_resp.json()["detail"]
